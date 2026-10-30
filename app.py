# -*- coding: utf-8 -*-
"""Habitat'Elles — prototype pédagogique de collecte et de réhabilitation.
Interface, modèles, ressources et exports sont regroupés dans ce fichier.
Les deux espaces sont propres à la session ; sauvegardes JSON explicites.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import html
import io
import json
import math
import re
import uuid
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from PIL import Image as PILImage, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image as PDFImage, KeepTogether

VERSION='1.0'
GREEN,CLAY,INK,PALE='#236757','#B85F43','#233D35','#F5F7F0'
CATEGORIES={'murs':'Murs','toiture':'Toiture','humidite':'Humidité','sols':'Sols','assainissement':'Assainissement'}
STATES=['Absent','Présent','Non observé','Information inconnue']
ORIGINS=['Observé pendant la visite','Déclaré par l’occupante','Mesure documentée']
EXTENTS=['Non renseignée','Localisée','Plusieurs points','Étendue']
HOUSING_TYPES=['Maison individuelle','Appartement','Chambre','Logement collectif','Autre']
SAMPLING=['Convenance / volontariat','Tirage aléatoire déclaré','Recensement du périmètre déclaré']
WEATHER=['Non renseigné','Temps sec','Pluie en cours','Pluie récente']
STAGES=['Signalé','Visite réalisée','Intervention proposée','Travaux réalisés','Résultat vérifié']
RESULTS=['Non vérifié','Amélioration observée','Problème persistant','Résolution constatée sur l’aspect suivi']
PRIORITIES=['Visite technique à prioriser','Réhabilitation à programmer','Entretien à organiser','Informations à compléter','Surveillance courante']
P_COLORS={PRIORITIES[0]:'#9A3D31',PRIORITIES[1]:'#B96B25',PRIORITIES[2]:'#6D7841',PRIORITIES[3]:'#6B7080',PRIORITIES[4]:GREEN}
LAWS=['Fixe','Uniforme','Triangulaire']
IMPORT_COLS=['rubrique','etat','origine','intensite','etendue','notes','preoccupant','motif_preoccupation']
REFS=[('OMS · logement et santé','https://www.who.int/publications/i/item/9789241550376'),
      ('OMS · humidité et moisissures','https://www.who.int/europe/publications/i/item/9789289041683'),
      ('Architecture connectée de Streamlit','https://docs.streamlit.io/develop/concepts/architecture/architecture'),
      ('Dépendances Community Cloud','https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/app-dependencies')]


def uid():return uuid.uuid4().hex[:12]


def fmt(value,digits=1):return f'{float(value):,.{digits}f}'.replace(',',' ').replace('.',',')


def checked_text(value,label,maximum=600,required=False):
    if not isinstance(value,str) or len(value)>maximum or (required and not value.strip()):raise ValueError(label+' : texte absent ou trop long.')
    return value.strip()


def number(value,label,low=0,high=1e9,integer=False):
    if isinstance(value,(bool,np.bool_)):raise ValueError(label+' : nombre attendu.')
    try:x=float(value.replace(',','.').replace(' ','')) if isinstance(value,str) else float(value)
    except (TypeError,ValueError):raise ValueError(label+' : nombre invalide.') from None
    if not math.isfinite(x) or not low<=x<=high or (integer and not x.is_integer()):raise ValueError(label+' : valeur hors limites ou non entière.')
    return int(x) if integer else x


def boolean(value,label):
    if type(value) is not bool:raise ValueError(label+' : booléen invalide.')
    return value


def iso(value,label,earliest=None,latest=None):
    try:d=date.fromisoformat(value)
    except (TypeError,ValueError):raise ValueError(label+' : date ISO AAAA-MM-JJ attendue.') from None
    if (earliest and d<earliest) or (latest and d>latest):raise ValueError(label+' : date hors période autorisée.')
    return d.isoformat()


def blank_aspect():
    return dict(state='Non observé',origin=ORIGINS[0],intensity=0,extent=EXTENTS[0],notes='',concern=False,reason='')


def new_project(demo=False):
    return dict(app="Habitat'Elles",version=VERSION,id=uid(),name='Enquête de démonstration' if demo else 'Mon enquête logement',demo=demo,
                start=date.today().isoformat(),sampling=SAMPLING[0],sampling_note='',neighborhoods=[],houses=[],visits=[],tasks=[])


def new_visit(p,house_id):
    return dict(id=uid(),house_id=house_id,seq=max([v['seq'] for v in p['visits']]+[0])+1,date=date.today().isoformat(),
                investigator='ENQ-01',weather=WEATHER[0],needs='',constraints='',aspects={k:blank_aspect() for k in CATEGORIES},photo_consent=False,photos=[])


def check_aspect(raw):
    if not isinstance(raw,dict) or set(raw)!=set(blank_aspect()):raise ValueError('Constat invalide ou champs non reconnus.')
    a=copy.deepcopy(raw)
    for key,choices in [('state',STATES),('origin',ORIGINS),('extent',EXTENTS)]:
        if a.get(key) not in choices:raise ValueError('Choix inconnu dans '+key)
    a['intensity']=number(a.get('intensity'),'Intensité',0,3,True)
    a['notes']=checked_text(a.get('notes'),'Notes du constat',600)
    a['reason']=checked_text(a.get('reason'),'Motif de préoccupation',500)
    a['concern']=boolean(a.get('concern'),'Signalement préoccupant')
    if a['state']=='Présent':
        if a['intensity']<1 or not a['notes']:raise ValueError('Un défaut présent nécessite une description et une intensité d’usage de 1 à 3.')
        if a['concern'] and not a['reason']:raise ValueError('Expliquer le signalement préoccupant.')
    elif a['intensity']!=0 or a['concern']:
        raise ValueError('Intensité et préoccupation concernent uniquement un défaut présent.')
    return a


def normalize_photo(raw):
    if not raw or len(raw)>5_000_000:raise ValueError('Photo limitée à 5 Mo.')
    try:
        with PILImage.open(io.BytesIO(raw)) as image:
            if image.format not in ['JPEG','PNG'] or image.width*image.height>20_000_000:raise ValueError('JPEG/PNG de 20 mégapixels maximum attendu.')
            clean=ImageOps.exif_transpose(image).convert('RGB');clean.thumbnail((1200,1200))
            data=None
            for quality in [80,65,50]:
                out=io.BytesIO();clean.save(out,format='JPEG',quality=quality,optimize=True);data=out.getvalue()
                if len(data)<=180_000:break
            if len(data)>180_000:
                clean.thumbnail((800,800));out=io.BytesIO();clean.save(out,format='JPEG',quality=60,optimize=True);data=out.getvalue()
            if len(data)>180_000:raise ValueError('Image encore trop volumineuse après réduction.')
    except (OSError,PILImage.DecompressionBombError):raise ValueError('Image illisible ou trop volumineuse.') from None
    return base64.b64encode(data).decode('ascii')


def validate(raw):
    if not isinstance(raw,dict) or raw.get('app')!="Habitat'Elles" or raw.get('version')!=VERSION:raise ValueError("Sauvegarde Habitat'Elles 1.0 attendue.")
    p=copy.deepcopy(raw);p['id']=checked_text(p.get('id'),'ID enquête',50,True);p['name']=checked_text(p.get('name'),'Nom enquête',120,True)
    if set(p)!=set(new_project()):raise ValueError('Champs de sauvegarde non reconnus.')
    p['demo']=boolean(p.get('demo'),'Mode fictif');p['start']=iso(p.get('start'),'Début enquête',latest=date.today())
    if p.get('sampling') not in SAMPLING:raise ValueError('Méthode de recrutement inconnue.')
    p['sampling_note']=checked_text(p.get('sampling_note'),'Méthode détaillée',900)
    for key,limit in [('neighborhoods',40),('houses',200),('visits',500),('tasks',500)]:
        if not isinstance(p.get(key),list) or len(p[key])>limit:raise ValueError(f'{key} : maximum {limit} enregistrements.')
        if len({x.get('id') for x in p[key] if isinstance(x,dict)})!=len(p[key]):raise ValueError('Identifiants manquants ou dupliqués dans '+key)
    for q in p['neighborhoods']:
        if set(q)!=set(['id','name','lat','lon','source']):raise ValueError('Champs de quartier non reconnus.')
        q['id']=checked_text(q.get('id'),'ID quartier',50,True);q['name']=checked_text(q.get('name'),'Nom quartier',100,True)
        q['source']=checked_text(q.get('source'),'Source du point de quartier',500)
        if q.get('lat') is None and q.get('lon') is None:continue
        q['lat']=round(number(q.get('lat'),'Latitude',-90,90),2);q['lon']=round(number(q.get('lon'),'Longitude',-180,180),2)
        if not q['source']:raise ValueError('Documenter le point approximatif du quartier.')
    if len({q['name'].casefold() for q in p['neighborhoods']})!=len(p['neighborhoods']):raise ValueError('Noms de quartiers dupliqués.')
    qids={q['id'] for q in p['neighborhoods']};hids=set();aliases=set()
    for h in p['houses']:
        if set(h)!=set(['id','alias','neighborhood_id','type','occupants','consent']):raise ValueError('Champs de logement non reconnus.')
        h['id']=checked_text(h.get('id'),'ID logement',20,True)
        if not re.fullmatch(r'HE-\d{4}',h['id']):raise ValueError('Identifiant de logement HE-0001 attendu.')
        if h.get('neighborhood_id') not in qids or h.get('type') not in HOUSING_TYPES:raise ValueError('Quartier ou type de logement inconnu.')
        h['alias']=checked_text(h.get('alias'),'Repère pseudonymisé',80,True);h['occupants']=number(h.get('occupants'),'Occupants',0,100,True)
        h['consent']=boolean(h.get('consent'),'Accord enquête')
        if not h['consent']:raise ValueError('Un accord d’enquête est requis pour enregistrer la fiche.')
        pair=(h['neighborhood_id'],h['alias'].casefold())
        if pair in aliases:raise ValueError('Repère de logement déjà présent dans ce quartier.')
        aliases.add(pair);hids.add(h['id'])
    photos=0;seqs=set();vids={}
    for v in p['visits']:
        if set(v)!=set(['id','house_id','seq','date','investigator','weather','needs','constraints','aspects','photo_consent','photos']):raise ValueError('Champs de visite non reconnus.')
        v['id']=checked_text(v.get('id'),'ID visite',50,True)
        if v.get('house_id') not in hids:raise ValueError('Visite liée à un logement absent.')
        v['seq']=number(v.get('seq'),'Ordre de visite',1,1e6,True)
        if v['seq'] in seqs:raise ValueError('Ordre de visite dupliqué.')
        seqs.add(v['seq']);v['date']=iso(v.get('date'),'Date de visite',date.fromisoformat(p['start']),date.today())
        v['investigator']=checked_text(v.get('investigator'),'Code enquêteur',20,True)
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,20}',v['investigator']):raise ValueError('Utiliser un code d’enquêteur, sans nom ni espaces.')
        if v.get('weather') not in WEATHER:raise ValueError('Contexte météo inconnu.')
        for key in ['needs','constraints']:v[key]=checked_text(v.get(key),key,800)
        if not isinstance(v.get('aspects'),dict) or set(v['aspects'])!=set(CATEGORIES):raise ValueError('Les cinq rubriques du constat sont requises.')
        v['aspects']={k:check_aspect(a) for k,a in v['aspects'].items()}
        v['photo_consent']=boolean(v.get('photo_consent'),'Accord photographies')
        if not isinstance(v.get('photos'),list) or len(v['photos'])>2:raise ValueError('Deux photos maximum par visite.')
        if v['photos'] and not v['photo_consent']:raise ValueError('Accord photographique requis.')
        if len({im.get('id') for im in v['photos'] if isinstance(im,dict)})!=len(v['photos']):raise ValueError('ID photo dupliqué.')
        for im in v['photos']:
            if set(im)!=set(['id','caption','aspect','data']):raise ValueError('Champs de photo non reconnus.')
            im['id']=checked_text(im.get('id'),'ID photo',50,True);im['caption']=checked_text(im.get('caption'),'Légende photo',200,True)
            if im.get('aspect') not in CATEGORIES:raise ValueError('Rubrique de photo inconnue.')
            data=im.get('data')
            if not isinstance(data,str) or len(data)>250_000:raise ValueError('Photo encodée invalide ou trop grande.')
            try:
                image_bytes=base64.b64decode(data,validate=True)
                with PILImage.open(io.BytesIO(image_bytes)) as img:
                    if img.format!='JPEG' or img.width>1200 or img.height>1200 or img.getexif():raise ValueError('Photo JPEG normalisée sans EXIF requise.')
                    img.verify()
            except (ValueError,OSError) as ex:raise ValueError('Photo de sauvegarde invalide : '+str(ex)) from None
        photos+=len(v['photos']);vids[v['id']]=v
    if photos>40:raise ValueError('40 photos maximum par enquête.')
    for t in p['tasks']:
        if set(t)!=set(['id','house_id','visit_id','aspect','stage','action','hypothesis','checks','planned_date','work_date','verified_date','result','verification_note','verification_source','budget_include','cost','cost_source']):raise ValueError('Champs d’intervention non reconnus.')
        t['id']=checked_text(t.get('id'),'ID intervention',50,True)
        if t.get('house_id') not in hids or t.get('visit_id') not in vids or vids[t['visit_id']]['house_id']!=t['house_id']:raise ValueError('Intervention liée à une visite ou un logement absent.')
        if t.get('aspect') not in CATEGORIES or t.get('stage') not in STAGES or t.get('result') not in RESULTS:raise ValueError('État d’intervention inconnu.')
        for key,limit,req in [('action',600,True),('hypothesis',600,False),('checks',600,False),('verification_note',600,False),('verification_source',400,False),('cost_source',500,False)]:t[key]=checked_text(t.get(key),key,limit,req)
        for key in ['planned_date','work_date','verified_date']:
            if t.get(key) is not None:t[key]=iso(t[key],key,latest=date.today() if key!='planned_date' else None)
        if t['stage']=='Travaux réalisés' and not t.get('work_date'):raise ValueError('Indiquer la date des travaux réalisés.')
        if t['stage']=='Résultat vérifié':
            if t['result']=='Non vérifié' or not t.get('verified_date') or not t['verification_source'] or not t['verification_note']:raise ValueError('Un résultat vérifié nécessite date, résultat, note et source de vérification.')
            if t['verified_date']<vids[t['visit_id']]['date']:raise ValueError('Vérification antérieure à la visite de référence.')
            if t.get('work_date') and t['verified_date']<t['work_date']:raise ValueError('Vérification antérieure aux travaux.')
        elif t['result']!='Non vérifié' or t.get('verified_date'):raise ValueError('Un résultat n’est renseigné qu’au stade Résultat vérifié.')
        t['budget_include']=boolean(t.get('budget_include'),'Inclusion au scénario budget')
        cost=t.get('cost')
        if cost is not None:
            if not isinstance(cost,dict) or set(cost)!=set(['law','min','mode','max']) or cost.get('law') not in LAWS:raise ValueError('Loi de coût inconnue.')
            a,b,c=(number(cost.get(k),'Coût '+k,0,1e9) for k in ['min','mode','max'])
            if not a<=b<=c:raise ValueError('Coût : minimum ≤ mode/centrale ≤ maximum requis.')
            if cost['law']=='Fixe' and not a==b==c:raise ValueError('Un coût fixe exige trois valeurs identiques.')
            if cost['law']=='Uniforme' and not math.isclose(b,(a+c)/2,abs_tol=1e-6):raise ValueError('Un coût uniforme exige une centrale au milieu des bornes.')
            t['cost']=dict(law=cost['law'],min=a,mode=b,max=c)
            if not t['cost_source']:raise ValueError('Documenter l’estimation de coût, même fictive.')
    if len(json.dumps(p,ensure_ascii=False,allow_nan=False).encode())>12_000_000:raise ValueError('Enquête limitée à 12 Mo, photos comprises.')
    return p


def example():
    p=new_project(True);p['start']='2026-10-01';p['sampling']=SAMPLING[1];p['sampling_note']='Tirage aléatoire supposé uniquement pour la démonstration fictive ; aucune enquête réelle.'
    for i,name in enumerate(['Quartier Alpha fictif','Quartier Beta fictif','Quartier Gamma fictif']):
        p['neighborhoods'].append(dict(id=uid(),name=name,lat=round(3.86+i*.02,2),lon=round(11.50+i*.02,2),source='Points approximatifs de démonstration ; aucun logement réel localisé.'))
    for i in range(1,31):
        h=dict(id=f'HE-{i:04d}',alias=f'Logement fictif {i:02d}',neighborhood_id=p['neighborhoods'][(i-1)//10]['id'],type=HOUSING_TYPES[(i-1)%4],occupants=0 if i%7==0 else 2+i%5,consent=True);p['houses'].append(h)
        v=new_visit(p,h['id']);v.update(date=f'2026-10-{15+i%10:02d}',investigator='ENQ-01' if i%2 else 'ENQ-02',needs='Priorité exprimée fictive : traiter les infiltrations.' if i<=12 else 'Besoin à discuter lors du suivi.',constraints='Accès aux pièces à organiser.')
        for key in CATEGORIES:v['aspects'][key].update(state='Absent')
        roof=v['aspects']['toiture']
        if i<=12:roof.update(state='Présent',intensity=3 if i<=4 else 2 if i<=8 else 1,extent=EXTENTS[1],notes='Infiltration fictive signalée dans une pièce.',origin=ORIGINS[1] if i%3==0 else ORIGINS[0])
        elif i>20:roof['state']='Non observé' if i<=25 else 'Information inconnue'
        for key,condition,intensity in [('murs',i%4==0,2),('humidite',i<=10,2),('sols',i%5==0,1),('assainissement',i%6==0,2)]:
            if condition:v['aspects'][key].update(state='Présent',intensity=intensity,extent=EXTENTS[1],notes='Constat entièrement fictif pour tester la collecte.')
        if i==8:v['aspects']['murs'].update(concern=True,reason='Évolution préoccupante rapportée dans cet exemple ; examen technique à organiser.')
        if i>24:v['aspects']['humidite']['state']='Non observé'
        p['visits'].append(v)
    for i in [0,1]:
        v=new_visit(p,p['houses'][i]['id']);v['date']='2026-10-02';v['needs']='Ancienne visite fictive.'
        for a in v['aspects'].values():a['state']='Absent'
        p['visits'].append(v)
    for i in [0,1,5,7,11,17]:
        v=next(v for v in p['visits'] if v['house_id']==p['houses'][i]['id'] and v['date']!='2026-10-02')
        aspect='murs' if i==7 else 'toiture' if i<12 else 'assainissement'
        t=dict(id=uid(),house_id=v['house_id'],visit_id=v['id'],aspect=aspect,stage='Intervention proposée',action='Intervention fictive à préciser après examen.',hypothesis='Cause à vérifier, non diagnostiquée par le logiciel.',checks='Inspection et métrés à compléter.',planned_date='2026-11-15',work_date=None,verified_date=None,result='Non vérifié',verification_note='',verification_source='',budget_include=i<2,cost=None,cost_source='')
        if i==0:t.update(cost=dict(law='Uniforme',min=100000.,mode=150000.,max=200000.),cost_source='Fourchette fictive de démonstration.')
        if i==1:t.update(cost=dict(law='Fixe',min=100000.,mode=100000.,max=100000.),cost_source='Montant fixe entièrement fictif.')
        p['tasks'].append(t)
    return validate(p)


def neighborhood_name(p,nid):return next(q['name'] for q in p['neighborhoods'] if q['id']==nid)


def next_house_id(p):return f"HE-{max([int(h['id'][3:]) for h in p['houses']]+[0])+1:04d}"


def visit_signature(v):
    return json.dumps({k:v[k] for k in ['house_id','date','aspects','needs','constraints']},ensure_ascii=False,sort_keys=True)


def put_visit(p,v):
    q=copy.deepcopy(p)
    if any(x['id']!=v['id'] and visit_signature(x)==visit_signature(v) for x in q['visits']):raise ValueError('Une visite identique existe déjà pour ce logement à cette date ; corriger la visite existante.')
    q['visits']=[x for x in q['visits'] if x['id']!=v['id']]+[v]
    return validate(q)


def latest_visits(p,start=None,end=None,area=None):
    hids={h['id'] for h in p['houses'] if area is None or h['neighborhood_id']==area};chosen={}
    for v in p['visits']:
        if v['house_id'] not in hids or (start and v['date']<start) or (end and v['date']>end):continue
        old=chosen.get(v['house_id'])
        if old is None or (v['date'],v['seq'])>(old['date'],old['seq']):chosen[v['house_id']]=v
    return chosen


def unresolved_concerns(p,house_id,until=None):
    flags=[]
    for v in p['visits']:
        if v['house_id']!=house_id or (until and v['date']>until):continue
        for key,a in v['aspects'].items():
            if not a['concern']:continue
            done=any(t['visit_id']==v['id'] and t['aspect']==key and t['stage']=='Résultat vérifié' and t['result']==RESULTS[-1] and (not until or t['verified_date']<=until) for t in p['tasks'])
            if not done:flags.append(dict(visit_id=v['id'],date=v['date'],aspect=key,reason=a['reason']))
    return flags


def priority(p,hid,visit=None,cutoff=2,until=None):
    alerts=unresolved_concerns(p,hid,until)
    if visit is None:visit=latest_visits(p,end=until).get(hid)
    if not visit:return dict(label=PRIORITIES[3],reasons=['Aucune visite disponible dans le périmètre retenu.'],evaluated=0,missing=5)
    a=visit['aspects'];evaluated=sum(x['state'] in STATES[:2] for x in a.values());missing=5-evaluated
    if alerts:
        label=PRIORITIES[0];reasons=[CATEGORIES[x['aspect']]+' · '+x['date']+' : '+x['reason'] for x in alerts]
    elif any(x['state']=='Présent' and x['intensity']>=cutoff for x in a.values()):
        label=PRIORITIES[1];reasons=[CATEGORIES[k]+f" : impact d’usage {x['intensity']}/3." for k,x in a.items() if x['state']=='Présent' and x['intensity']>=cutoff]
    elif any(x['state']=='Présent' for x in a.values()):label=PRIORITIES[2];reasons=['Défaut présent, sans atteindre le seuil d’usage du classement de réhabilitation.']
    elif missing:label=PRIORITIES[3];reasons=['Aucun défaut présent consigné, mais des rubriques ne sont pas évaluées.']
    else:label=PRIORITIES[4];reasons=['Aucun défaut consigné dans les cinq rubriques évaluées à cette visite ; ce constat ne qualifie pas tout le logement.']
    if missing:reasons.append(f'{missing} rubrique(s) à compléter ; les données manquantes ne valent pas absence de défaut.')
    return dict(label=label,reasons=reasons,evaluated=evaluated,missing=missing)


def wilson(k,n):
    if n<=0:return None,None
    z=1.95996398454;p=k/n;den=1+z*z/n;center=(p+z*z/(2*n))/den;half=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/den
    return max(0.,center-half),min(1.,center+half)


def rates(p,start=None,end=None,area=None):
    chosen=latest_visits(p,start,end,area);rows=[]
    for key,label in CATEGORIES.items():
        aspects=[v['aspects'][key] for v in chosen.values()];counts={s:sum(a['state']==s for a in aspects) for s in STATES}
        n=counts['Absent']+counts['Présent'];lo,hi=wilson(counts['Présent'],n)
        rows.append(dict(rubrique=label,logements_visites=len(chosen),presents=counts['Présent'],absents=counts['Absent'],evalues=n,
                         non_observes=counts['Non observé'],inconnus=counts['Information inconnue'],pourcentage=100*counts['Présent']/n if n else None,
                         IC95_bas=100*lo if lo is not None else None,IC95_haut=100*hi if hi is not None else None,
                         presents_observes=sum(a['state']=='Présent' and a['origin']==ORIGINS[0] for a in aspects),
                         presents_declares=sum(a['state']=='Présent' and a['origin']==ORIGINS[1] for a in aspects),
                         presents_mesures=sum(a['state']=='Présent' and a['origin']==ORIGINS[2] for a in aspects)))
    return pd.DataFrame(rows)


def housing_table(p,start=None,end=None,area=None,cutoff=2):
    chosen=latest_visits(p,start,end,area);rows=[]
    for h in p['houses']:
        if area and h['neighborhood_id']!=area:continue
        v=chosen.get(h['id']);pr=priority(p,h['id'],v,cutoff,end) if v else dict(label=PRIORITIES[3],evaluated=0,missing=5,reasons=['Aucune visite dans cette période.'])
        rows.append(dict(logement=h['id'],quartier=neighborhood_name(p,h['neighborhood_id']),type=h['type'],derniere_visite=v['date'] if v else '',priorite=pr['label'],rubriques_evaluees=pr['evaluated'],rubriques_manquantes=pr['missing'],occupants=h['occupants'] if h['occupants'] else None))
    return pd.DataFrame(rows,columns=['logement','quartier','type','derniere_visite','priorite','rubriques_evaluees','rubriques_manquantes','occupants'])


def neighborhood_table(p,start=None,end=None,area=None):
    chosen=latest_visits(p,start,end,area);rows=[]
    for q in p['neighborhoods']:
        if area and q['id']!=area:continue
        houses=[h for h in p['houses'] if h['neighborhood_id']==q['id']]
        rows.append(dict(quartier=q['name'],recenses=len(houses),visites=sum(h['id'] in chosen for h in houses),lat=q['lat'],lon=q['lon']))
    return pd.DataFrame(rows,columns=['quartier','recenses','visites','lat','lon'])


def map_points(p,start=None,end=None,area=None):
    q=neighborhood_table(p,start,end,area)
    return q[(q.visites>=5)&q.lat.notna()&q.lon.notna()].copy()


def aspect_rows(p):
    rows=[]
    for v in p['visits']:
        for key,a in v['aspects'].items():rows.append(dict(logement=v['house_id'],visite=v['id'],date=v['date'],enqueteur=v['investigator'],rubrique=key,etat=a['state'],origine=a['origin'],intensite=a['intensity'],etendue=a['extent'],notes=a['notes'],preoccupant=a['concern'],motif_preoccupation=a['reason']))
    return pd.DataFrame(rows,columns=['logement','visite','date','enqueteur','rubrique','etat','origine','intensite','etendue','notes','preoccupant','motif_preoccupation'])


def import_aspects(raw,filename):
    if len(raw)>1_000_000:raise ValueError('Import limité à 1 Mo.')
    if filename.lower().endswith('.csv'):df=pd.read_csv(io.BytesIO(raw),sep=';',encoding='utf-8-sig',dtype=str,keep_default_na=False)
    elif filename.lower().endswith('.xlsx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist())>10_000_000:raise ValueError('Classeur décompressé trop volumineux.')
        df=pd.read_excel(io.BytesIO(raw),sheet_name='Constats',dtype=str,keep_default_na=False,engine='openpyxl')
    else:raise ValueError('CSV ou XLSX attendu.')
    if set(df.columns)!=set(IMPORT_COLS) or len(df)!=5 or set(df.rubrique)!=set(CATEGORIES):raise ValueError('Cinq lignes uniques et colonnes du modèle requises, une par rubrique.')
    aspects={}
    for row in df.to_dict('records'):
        flag=str(row['preoccupant']).strip().lower()
        if flag not in ['true','false','vrai','faux','1','0']:raise ValueError('preoccupant : true/false attendu.')
        aspects[row['rubrique']]=check_aspect(dict(state=row['etat'],origin=row['origine'],intensity=number(row['intensite'],'Intensité',0,3,True),extent=row['etendue'],notes=row['notes'],concern=flag in ['true','vrai','1'],reason=row['motif_preoccupation']))
    return aspects


def template_frame():
    return pd.DataFrame([dict(rubrique=k,etat='Non observé',origine=ORIGINS[0],intensite=0,etendue=EXTENTS[0],notes='',preoccupant=False,motif_preoccupation='') for k in CATEGORIES],columns=IMPORT_COLS)


def csv_bytes(df):
    safe=df.copy()
    for col in safe.select_dtypes(include=['object','string']).columns:
        safe[col]=safe[col].map(lambda x:"'"+x if isinstance(x,str) and x.lstrip().startswith(('=','+','-','@')) else x)
    return safe.to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')


def budget_scope(p):
    tasks=[t for t in p['tasks'] if t['budget_include'] and t['stage'] not in STAGES[-2:]]
    return [t for t in tasks if t['cost'] is not None],sum(t['cost'] is None for t in tasks)


@st.cache_data(ttl=1800,max_entries=4,show_spinner=False)
def budget_simulation(costs,count=10000,seed=42):
    count=number(count,'Tirages',1000,100000,True);seed=number(seed,'Graine',0,2**32-1,True);rng=np.random.default_rng(seed);total=np.zeros(count)
    if not costs:raise ValueError('Aucune intervention chiffrée sélectionnée dans le scénario.')
    for law,a,b,c in costs:
        if law not in LAWS or not 0<=a<=b<=c or c>1e9:raise ValueError('Hypothèse de coût invalide.')
        if a==c or law=='Fixe':total+=b
        elif law=='Uniforme':total+=rng.uniform(a,c,count)
        else:total+=rng.triangular(a,b,c,count)
    return total


def backup_bytes(p):return json.dumps(validate(p),ensure_ascii=False,indent=2,allow_nan=False).encode('utf-8')


def restore(raw):
    if len(raw)>12_000_000:raise ValueError('Sauvegarde limitée à 12 Mo.')
    return validate(json.loads(raw.decode('utf-8-sig')))


def excel_bytes(p):
    wb=Workbook();wb.remove(wb.active)
    households=pd.DataFrame([dict(id=h['id'],repere=h['alias'],quartier=neighborhood_name(p,h['neighborhood_id']),type=h['type'],occupants=h['occupants'] if h['occupants'] else None,accord=h['consent']) for h in p['houses']],columns=['id','repere','quartier','type','occupants','accord'])
    visits=pd.DataFrame([dict(id=v['id'],logement=v['house_id'],date=v['date'],ordre=v['seq'],enqueteur=v['investigator'],meteo=v['weather'],besoin_exprime=v['needs'],contraintes=v['constraints'],accord_photos=v['photo_consent'],nombre_photos=len(v['photos'])) for v in p['visits']],columns=['id','logement','date','ordre','enqueteur','meteo','besoin_exprime','contraintes','accord_photos','nombre_photos'])
    tasks=pd.DataFrame([dict(id=t['id'],logement=t['house_id'],visite=t['visit_id'],rubrique=CATEGORIES[t['aspect']],action=t['action'],cause_supposee=t['hypothesis'],verification_necessaire=t['checks'],stade=t['stage'],date_cible=t['planned_date'],date_travaux=t['work_date'],date_verification=t['verified_date'],resultat=t['result'],note_verification=t['verification_note'],source_verification=t['verification_source'],scenario_budget=t['budget_include'],loi_cout=t['cost']['law'] if t['cost'] else '',cout_min=t['cost']['min'] if t['cost'] else None,cout_central=t['cost']['mode'] if t['cost'] else None,cout_max=t['cost']['max'] if t['cost'] else None,source_cout=t['cost_source']) for t in p['tasks']])
    tables={'Logements':households,'Visites':visits,'Constats':aspect_rows(p),'Interventions':tasks,'Statistiques':rates(p),'Quartiers':neighborhood_table(p),
            'Methode':pd.DataFrame([('Enquête',p['name']),('Fictif',str(p['demo'])),('Début',p['start']),('Recrutement',p['sampling']),('Note méthode',p['sampling_note']),('Lecture','Statistiques sur la dernière visite de chaque logement ; inconnus exclus du dénominateur évalué.'),('IC95','Calcul mathématique sous modèle binomial ; ne corrige pas les biais de convenance et ne représente pas une incertitude de recensement.'),('Photos','Images non incluses dans Excel ; sauvegarde JSON dédiée.')],columns=['champ','valeur'])}
    for name,frame in tables.items():
        ws=wb.create_sheet(name);ws.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):ws.append([x.item() if isinstance(x,np.generic) else None if pd.isna(x) else x for x in row])
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='236757')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=27
    out=io.BytesIO();wb.save(out);return out.getvalue()


def pdf_bytes(p,house_id=None,include_photos=False):
    if 'HabitatVera' not in pdfmetrics.getRegisteredFontNames():
        f=Path(rl_config.__file__).parent/'fonts';pdfmetrics.registerFont(TTFont('HabitatVera',str(f/'Vera.ttf')));pdfmetrics.registerFont(TTFont('HabitatBold',str(f/'VeraBd.ttf')))
    s=getSampleStyleSheet();s.add(ParagraphStyle('HB',fontName='HabitatVera',fontSize=9,leading=13,spaceAfter=7,splitLongWords=True));s.add(ParagraphStyle('HT',parent=s['HB'],fontName='HabitatBold',fontSize=22,leading=28,textColor=colors.HexColor(GREEN),spaceAfter=13));s.add(ParagraphStyle('HH',parent=s['HB'],fontName='HabitatBold',fontSize=12,leading=17,spaceBefore=12,keepWithNext=True));s.add(ParagraphStyle('HC',parent=s['HB'],fontSize=8,leading=11,spaceAfter=0))
    s.add(ParagraphStyle('HI',parent=s['HB'],keepWithNext=True))
    def par(t,style='HB'):return Paragraph(html.escape(str(t)).replace('\n','<br/>'),s[style])
    def table(head,rows,widths):
        t=Table([[par(x,'HC') for x in head]]+[[par(x,'HC') for x in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE9DC')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor(PALE)]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));return t
    output=io.BytesIO();doc=SimpleDocTemplate(output,pagesize=(21*cm,29.7*cm),leftMargin=1.7*cm,rightMargin=1.7*cm,topMargin=1.6*cm,bottomMargin=1.7*cm)
    story=[par("Habitat’Elles",'HT'),par(('EXEMPLE FICTIF · ' if p['demo'] else '')+p['name'])]
    if house_id is None:
        chosen=latest_visits(p);stats=rates(p);ht=housing_table(p)
        story += [par('Rapport agrégé de l’enquête','HH'),par('Sans fiches individuelles, photos, codes de logement ni points cartographiques. Les résultats décrivent les données du registre, sans diagnostic de sécurité des logements.'),
                  par(f"Début : {p['start']}. Logements recensés : {len(p['houses'])}. Logements avec visite : {len(chosen)}. Visites enregistrées : {len(p['visits'])}. Une seule visite par logement est retenue pour les fréquences : la plus récente, puis l’ordre d’enregistrement à date égale."),
                  par('Recrutement déclaré : '+p['sampling']+'. Une sélection par convenance ne permet pas de généraliser automatiquement au quartier.'),
                  par('Défauts consignés et données manquantes','HH'),
                  table(['Rubrique','Présents / évalués','Non observés / inconnus','Fréquence'],[[r.rubrique,f'{r.presents} / {r.evalues}',f'{r.non_observes} / {r.inconnus}',fmt(r.pourcentage)+' %' if pd.notna(r.pourcentage) else 'Non définie'] for r in stats.itertuples()],[4.2*cm,4.3*cm,5.1*cm,4*cm]),
                  par('Évalués = présent + absent. Non observé et information inconnue sont exclus du dénominateur. Les origines observées, déclarées ou mesurées sont détaillées dans les tableaux de l’application et du classeur. Les défauts de plusieurs rubriques peuvent coexister dans un même logement.'),
                  par('Priorités de traitement proposées','HH')]
        counts=ht.priorite.value_counts() if len(ht) else pd.Series(dtype=int)
        story.append(table(['Orientation','Logements'],[[label,int(counts.get(label,0))] for label in PRIORITIES],[13*cm,4.6*cm]))
        qs=neighborhood_table(p);public=qs[qs.recenses>=5];small=qs[qs.recenses<5]
        story.append(par('Répartition par quartier','HH'))
        rows=[[r.quartier,r.recenses,r.visites] for r in public.itertuples()]
        if len(small):rows.append(['Autres quartiers (groupes de moins de 5 logements)',int(small.recenses.sum()),int(small.visites.sum())])
        if rows:story.append(table(['Quartier ou regroupement','Recensés','Avec visite'],rows,[10*cm,3.8*cm,3.8*cm]))
        known,unknown=budget_scope(p)
        story += [par('Suivi et scénario budgétaire','HH'),par(f"Interventions enregistrées : {len(p['tasks'])}. Scénario sélectionné : {len(known)} interventions futures chiffrées et {unknown} non chiffrées. Les travaux réalisés sont exclus de ce scénario de dépenses futures."),
                  par('Les montants saisis sont des estimations à documenter, pas des prix de marché fournis par le logiciel. Les résultats Monte-Carlo s’exportent séparément depuis le suivi.')]
    else:
        h=next((h for h in p['houses'] if h['id']==house_id),None)
        if h is None:raise ValueError('Logement absent du registre.')
        pr=priority(p,h['id']);visits=sorted([v for v in p['visits'] if v['house_id']==h['id']],key=lambda v:(v['date'],v['seq']),reverse=True)
        story += [par('Fiche de suivi détaillée · usage réservé','HH'),par(f"{h['id']} · {neighborhood_name(p,h['neighborhood_id'])} · {h['type']}. Occupants : "+(str(h['occupants']) if h['occupants'] else 'non renseigné')),
                  par('Orientation : '+pr['label'])]+[par('• '+r) for r in pr['reasons']]
        for v in visits[:3]:
            story += [par('Visite du '+v['date'],'HH'),par('Code enquêteur : '+v['investigator']+' ; contexte : '+v['weather'],'HI'),
                      table(['Rubrique','Constat / origine','Impact d’usage','Description'],[[CATEGORIES[k],a['state']+' · '+a['origin'],str(a['intensity'])+'/3' if a['state']=='Présent' else 'Non applicable',a['notes']] for k,a in v['aspects'].items()],[3.2*cm,5*cm,3.2*cm,6.2*cm]),
                      par('Besoin exprimé : '+(v['needs'] or 'Non renseigné')),par('Contraintes pratiques : '+(v['constraints'] or 'Non renseignées'))]
            if include_photos and v['photo_consent']:
                for photo in v['photos']:
                    data=base64.b64decode(photo['data']);image=PILImage.open(io.BytesIO(data));factor=min(16*cm/image.width,6*cm/image.height)
                    story.append(KeepTogether([par('Photo autorisée · '+CATEGORIES[photo['aspect']]+' · '+photo['caption']),PDFImage(io.BytesIO(data),width=image.width*factor,height=image.height*factor)]))
        story.append(par('La fiche reprend les trois dernières visites au maximum. Le registre complet reste dans la sauvegarde JSON. Les photos sont incluses uniquement sur demande et si l’accord photographique a été enregistré.'))
        for t in [t for t in p['tasks'] if t['house_id']==h['id']]:
            story += [par(CATEGORIES[t['aspect']]+' · '+t['stage'],'HH'),par('Action : '+t['action']),par('Cause supposée : '+(t['hypothesis'] or 'À examiner')+' ; vérifications : '+(t['checks'] or 'À préciser')),par('Résultat : '+t['result']+' ; '+t['verification_note'])]
    story += [par('Lecture et limites','HH'),par('L’orientation est une proposition de traitement à examiner. Un signalement préoccupant non clôturé par une vérification documentée reste prioritaire, même si une visite plus récente ne réévalue pas la rubrique. Le classement ne repose pas sur le revenu, le statut familial ou le nombre d’occupants. Une photo ou un seul indice ne constitue pas un diagnostic structurel.'),
              par('Le registre distingue constats de visite, déclarations et mesures. La qualité d’un échantillon ne découle pas du nombre de logements seul. Les distributions et intervalles d’exercice ne corrigent ni les biais de sélection ni les erreurs de collecte. Travaux réalisés et résultat vérifié sont deux étapes distinctes.'),
              par('Références de méthode','HH')]+[par(title+' : '+url) for title,url in REFS[:2]]
    def footer(canvas,doc):
        canvas.setFillColor(colors.HexColor(GREEN));canvas.rect(0,29.7*cm-.25*cm,21*cm,.25*cm,fill=1,stroke=0);canvas.setFillColor(colors.HexColor(INK));canvas.setFont('HabitatVera',8);canvas.drawString(1.7*cm,1*cm,"Habitat’Elles · "+('FICTIF · ' if p['demo'] else '')+('Synthèse agrégée' if house_id is None else 'Fiche détaillée'));canvas.drawRightString(19.3*cm,1*cm,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return output.getvalue()


CSS='''<style>
.stApp{background:#F5F7F0;color:#233D35}[data-testid="stMainBlockContainer"]{max-width:1460px;padding-top:4.8rem;padding-bottom:3rem}
[data-testid="stSidebar"]{background:#FBFCF6;border-right:1px solid #DCE5D6;color:#24483D}[data-testid="stSidebarNav"] a span{color:#24483D}[data-testid="stSidebarNav"] a:hover{background:#E2ECDF}
h1,h2,h3{color:#214E40;letter-spacing:-.035em}h1{font-size:2.25rem!important}p,label{line-height:1.65}
.mast{display:flex;align-items:center;justify-content:space-between;gap:12px;flex-wrap:wrap;margin-bottom:25px}.brand{font-weight:900;font-size:1.4rem;letter-spacing:-.055em;color:#236757}.brand span{color:#B85F43}.tag{background:#E1ECDC;color:#2B5544;padding:7px 12px;border-radius:30px;font-size:.68rem;letter-spacing:.08em;font-weight:750}
.hero{display:grid;grid-template-columns:1.15fr 1fr;gap:28px;align-items:center;background:#FFFDF6;border:1px solid #D9E2D3;border-radius:27px;padding:32px 36px;margin-bottom:25px;box-shadow:0 12px 35px #24473708}.hero h2{font-size:2.65rem;line-height:1.24;margin:12px 0 15px;max-width:550px}.hero p{color:#5A6C5E;font-size:.93rem;max-width:560px}.hero img{width:100%;max-height:255px;background:#DCE9D7;border-radius:19px}.eyebrow{font-size:.69rem;letter-spacing:.13em;color:#66785D;font-weight:800}
.kpis{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px;margin:15px 0 24px}.kpi{border:1px solid #DCE5D5;background:#FFFDF7;border-radius:18px;padding:21px;min-width:0}.kpi:nth-child(2){background:#236757;color:#FFFDF6;border-color:#236757}.kpi:nth-child(3){background:#F5DFD0;border-color:#E9CEB9}.kpi:nth-child(4){background:#DFE8D4}.kpi label{display:block;line-height:1.45;font-size:.78rem;opacity:.85}.kpi strong{display:block;font-size:1.85rem;font-weight:850;letter-spacing:-.04em;line-height:1.3;margin:9px 0;overflow-wrap:anywhere}.kpi small{display:block;line-height:1.55;font-size:.72rem;opacity:.85}
.note{background:#E4EEDC;border-left:4px solid #67875D;padding:15px 19px;border-radius:0 12px 12px 0;color:#2B503E;line-height:1.65;margin:12px 0 22px}.housecard{border:1px solid #D7E0CE;border-radius:17px;padding:21px;background:#FFFDF7;margin-bottom:10px}.housecard h3{font-size:1.15rem!important;margin:5px 0}.housecard small{font-size:.75rem;color:#61735F}.pill{display:inline-block;font-size:.72rem;padding:5px 9px;border-radius:6px;background:#E9EEDC;font-weight:700;margin-top:9px;white-space:normal}.sidebrand{font-size:1.7rem;font-weight:850;color:#236757;margin:15px 0 10px;letter-spacing:-.06em}.sidenote{color:#6D7B64;font-size:.78rem;line-height:1.7;margin-bottom:20px}.foot{font-size:.75rem;line-height:1.6;color:#667660;padding-top:25px;border-top:1px solid #DCE5D3;margin-top:34px}
[data-testid="stForm"]{background:#FFFDF7;border-color:#D5DFCB;border-radius:18px;padding:22px}[data-testid="stButton"] button[kind="primary"],[data-testid="stFormSubmitButton"] button[kind="primary"]{background:#236757;color:#FFFFFF;border-color:#236757;border-radius:11px;min-height:44px}button p{white-space:normal;overflow-wrap:anywhere}a{overflow-wrap:anywhere}
@media(max-width:950px){.kpis{grid-template-columns:repeat(2,minmax(0,1fr))}.hero{grid-template-columns:1fr .85fr;padding:26px}.hero h2{font-size:2.1rem}}
@media(max-width:600px){[data-testid="stMainBlockContainer"]{padding:4.5rem 1rem 2rem}.hero{grid-template-columns:1fr;padding:23px}.hero h2{font-size:1.95rem}.hero img{max-height:185px}.kpi{padding:15px}.kpi strong{font-size:1.4rem}.kpis{gap:10px}.brand{font-size:1.25rem}h1{font-size:1.9rem!important}.tag{font-size:.61rem}}
</style>'''


def hero_art():
    svg='''<svg xmlns="http://www.w3.org/2000/svg" width="480" height="290" viewBox="0 0 480 290"><rect width="480" height="290" fill="#DCE9D7"/><circle cx="410" cy="44" r="57" fill="#B7D3B8"/><path d="M0 255c100-45 150 15 240-10s160-25 240 3v42H0Z" fill="#C2D7BB"/><path d="M113 133 238 54 363 133v120H113Z" fill="#FFF9EB"/><path d="m95 139 143-93 143 93-17 17-126-80-126 80Z" fill="#B85F43"/><path d="M238 76v177h125V154Z" fill="#E8DFCA"/><rect x="146" y="156" width="61" height="59" rx="5" fill="#A5C6BA"/><path d="M176 157v58m-29-29h60" stroke="#FFF9EB" stroke-width="4"/><rect x="268" y="166" width="50" height="87" rx="4" fill="#65937C"/><circle cx="304" cy="209" r="3" fill="#F8E6C8"/><path d="m228 144-9 18 12 15-10 22 10 16" fill="none" stroke="#A57059" stroke-width="3"/><circle cx="76" cy="91" r="29" fill="#236757"/><path d="m63 91 9 9 18-20" fill="none" stroke="#F5F7ED" stroke-width="4"/><rect x="332" y="183" width="113" height="59" rx="10" fill="#FFFBF1"/><path d="M347 204h66m-66 13h47m-47 13h57" stroke="#84A085" stroke-width="5" stroke-linecap="round"/><path d="M63 242v-57m-18 24 18-24 18 24" fill="none" stroke="#7A9A76" stroke-width="5"/></svg>'''
    return 'data:image/svg+xml;base64,'+base64.b64encode(svg.encode()).decode()


def plot(fig,alt,key=None):
    fig.update_layout(template='plotly_white',font=dict(family='Arial',size=12,color=INK),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',margin=dict(l=12,r=12,t=65,b=45),height=380,colorway=[GREEN,CLAY,'#829D70','#C3C7B5','#8E93A6'],legend=dict(orientation='h',x=0,y=1.03,yanchor='bottom'))
    st.plotly_chart(fig,width='stretch',key=key,alt=alt,config={'displaylogo':False,'scrollZoom':False})


def note(t):st.html('<div class="note">'+t+'</div>')


def kpis(items):
    st.html('<div class="kpis">'+''.join('<div class="kpi"><label>'+html.escape(str(label))+'</label><strong>'+html.escape(str(value))+'</strong><small>'+html.escape(str(detail))+'</small></div>' for label,value,detail in items)+'</div>')


def title(t,subtitle):st.title(t);st.caption(subtitle)


def current():return st.session_state.demo_project if st.session_state.demo else st.session_state.real_project


def commit(p):
    q=validate(p)
    if st.session_state.demo:st.session_state.demo_project=q
    else:st.session_state.real_project=q
    st.session_state.rev+=1;st.session_state.budget_result=None;st.session_state.flash='Enregistrement effectué dans cette session.';st.rerun()


def choose_house(p,key):
    if not p['houses']:
        st.info('Commencez par enregistrer un quartier puis un logement.');st.page_link(st.session_state.pages['logements'],label='Créer une fiche logement',icon=':material/home:');return None
    ids=[h['id'] for h in p['houses']];active=st.session_state.get('active_house');index=ids.index(active) if active in ids else 0
    hid=st.selectbox('Logement à consulter',ids,index=index,format_func=lambda i:i+' · '+neighborhood_name(p,next(h['neighborhood_id'] for h in p['houses'] if h['id']==i)),key=key+'_'+p['id'])
    st.session_state.active_house=hid
    return next(h for h in p['houses'] if h['id']==hid)


def go_house(hid,page):
    st.session_state.active_house=hid;st.switch_page(st.session_state.pages[page])


def dashboard():
    p=current();chosen=latest_visits(p);ht=housing_table(p)
    st.html('<div class="hero"><div><div class="eyebrow">CAMPAGNE DE COLLECTE / HABITAT & RÉHABILITATION</div><h2>Des constats.<br>Des actions à suivre.</h2><p>Recensez les logements, écoutez les besoins exprimés et organisez les étapes de réhabilitation.</p><span class="tag">'+('DÉMONSTRATION FICTIVE' if p['demo'] else 'ESPACE PERSONNEL')+'</span></div><img src="'+hero_art()+'" alt="Illustration d’un logement et de sa fiche de visite ; dessin sans valeur de preuve"/></div>')
    urgent=int((ht.priorite==PRIORITIES[0]).sum()) if len(ht) else 0
    kpis([('Logements recensés',len(p['houses']),'Un identifiant par logement'),('Avec une visite',len(chosen),f"{len(p['visits'])} visites dans l’historique"),('Visites à prioriser',urgent,'Signalements à faire examiner'),('Interventions ouvertes',sum(t['stage'] not in STAGES[-2:] for t in p['tasks']),'Travaux / vérifications à suivre')])
    if not p['houses']:
        note('Votre espace est vide. Créez un quartier et un logement, ou ouvrez l’exemple fictif sans modifier vos données.')
        st.page_link(st.session_state.pages['logements'],label='Créer mes premières fiches',icon=':material/add_home:');return
    with st.expander('Période et quartier de l’analyse'):
        with st.form('analysis_filters_'+str(st.session_state.rev)):
            saved_start,saved_end,saved_area=st.session_state.filters.get(p['id'],(None,None,None))
            a,b=st.columns(2);start=a.date_input('Visites à partir du',date.fromisoformat(saved_start or p['start']),max_value=date.today());end=b.date_input('Visites jusqu’au',date.fromisoformat(saved_end) if saved_end else date.today(),max_value=date.today())
            areas=['Tous']+[q['id'] for q in p['neighborhoods']]
            area=st.selectbox('Quartier de l’analyse',areas,index=areas.index(saved_area) if saved_area in areas else 0,format_func=lambda x:'Tous les quartiers' if x=='Tous' else neighborhood_name(p,x))
            if st.form_submit_button('Appliquer les filtres'):
                if start>end:raise ValueError('La date de début doit précéder la date de fin.')
                st.session_state.filters[p['id']]=(start.isoformat(),end.isoformat(),None if area=='Tous' else area)
    start,end,area=st.session_state.filters.get(p['id'],(None,None,None))
    filtered=latest_visits(p,start,end,area);stats=rates(p,start,end,area);filtered_ht=housing_table(p,start,end,area)
    st.caption(f"Analyse : {len(filtered)} logements avec visite, du {start or p['start']} au {end or date.today().isoformat()} · dernière visite de chaque logement dans cette période. Les indicateurs du bandeau décrivent le registre entier.")
    a,b=st.columns([1.3,1])
    with a:
        st.subheader('Ce qui a été consigné')
        fig=go.Figure()
        for col,label,color in [('presents','Présent',CLAY),('absents','Absent',GREEN),('non_observes','Non observé','#B9C5A9'),('inconnus','Inconnu','#D9DDD0')]:fig.add_bar(y=stats.rubrique,x=stats[col],name=label,orientation='h',marker_color=color)
        fig.update_layout(barmode='stack');fig.update_xaxes(title='Logements distincts');plot(fig,'Nombre de logements par rubrique et statut, données manquantes visibles')
    with b:
        st.subheader('Orientations proposées')
        counts=filtered_ht[filtered_ht.derniere_visite!=''].priorite.value_counts() if len(filtered_ht) else pd.Series(dtype=int)
        labels=[label for label in PRIORITIES if counts.get(label,0)]
        if labels:
            fig=go.Figure(go.Pie(labels=labels,values=[counts[label] for label in labels],hole=.70,marker_colors=[P_COLORS[x] for x in labels],textinfo='value',hovertemplate='%{label}<br>%{value} logement(s)<extra></extra>'));fig.update_layout(showlegend=False);plot(fig,'Répartition des orientations des logements visités')
            for label in labels:st.caption(f'{label} : {counts[label]}')
        else:st.caption('Aucune visite dans le périmètre retenu.')
    st.subheader('Des dénominateurs visibles')
    columns=['rubrique','presents','evalues','non_observes','inconnus','pourcentage']
    display=stats[columns].round({'pourcentage':1}).rename(columns={'rubrique':'Rubrique','presents':'Présents','evalues':'Évalués','non_observes':'Non observés','inconnus':'Inconnus','pourcentage':'Présents / évalués (%)'})
    st.dataframe(display,hide_index=True,alt='Fréquences des défauts parmi les rubriques évaluées, avec dénominateurs et données manquantes')
    note('<b>Présents / évalués</b> : seuls « présent » et « absent » entrent au dénominateur. Une même habitation peut présenter plusieurs défauts. Les déclarations et observations restent identifiées ; aucune relation causale n’est déduite.')
    with st.expander('Statistiques et probabilités : lire l’incertitude'):
        st.markdown('**Origine des signalements présents**')
        origins=stats[['rubrique','presents_observes','presents_declares','presents_mesures']].rename(columns={'rubrique':'Rubrique','presents_observes':'Observés','presents_declares':'Déclarés','presents_mesures':'Mesurés'})
        st.dataframe(origins,hide_index=True,alt='Origine des constats présents : observation, déclaration ou mesure documentée')
        st.write('Recrutement déclaré : **'+p['sampling']+'**. '+p['sampling_note'])
        show=st.checkbox('Afficher les intervalles Wilson comme exercice sous modèle binomial',value=p['demo'] and p['sampling']==SAMPLING[1])
        if show:
            st.dataframe(stats[['rubrique','presents','evalues','pourcentage','IC95_bas','IC95_haut']],hide_index=True,alt='Intervalles Wilson conditionnels au modèle binomial, non correction des biais d’enquête')
            st.caption('Ces intervalles supposent un modèle d’observations suffisamment indépendantes et représentatives. Ils ne corrigent pas un recrutement par convenance, la non-réponse ou les erreurs de mesure. Ils ne représentent pas une incertitude de recensement exhaustif. Si aucun aspect n’est évalué, fréquence et intervalle restent non définis.')
    with st.expander('Carte des quartiers et répartition territoriale'):
        q=neighborhood_table(p,start,end,area);st.dataframe(q[['quartier','recenses','visites']],hide_index=True,alt='Logements recensés et visités par quartier')
        st.caption('La carte ne reçoit que des points de quartier arrondis à 0,01° et des effectifs ; jamais une position de logement. Seuls les quartiers avec au moins cinq logements visités apparaissent. Ce regroupement réduit la précision sans garantir à lui seul l’anonymat.')
        mp=map_points(p,start,end,area)
        if len(mp):
            st.markdown('**Repères de quartiers, sans fond externe**')
            labels=[str(i+1) for i in range(len(mp))]
            fig=go.Figure(go.Scatter(x=mp.lon,y=mp.lat,mode='markers+text',text=labels,textposition='middle center',textfont=dict(color='white'),marker=dict(size=28,color=GREEN),customdata=mp[['quartier','visites']].to_numpy(),hovertemplate='%{customdata[0]}<br>%{customdata[1]} logements visités<extra></extra>'))
            fig.update_xaxes(title='Longitude approximative');fig.update_yaxes(title='Latitude approximative',scaleanchor='x',scaleratio=1);plot(fig,'Repères approximatifs des quartiers, indépendants des fonds cartographiques en ligne')
            for i,row in enumerate(mp.itertuples(),1):st.caption(f'{i} · {row.quartier} · {row.visites} logements visités')
        if st.checkbox('Afficher la carte agrégée (connexion aux fonds cartographiques)'):
            if len(mp):
                span=max(float(mp.lon.max()-mp.lon.min()),float(mp.lat.max()-mp.lat.min()),.01)
                zoom=int(max(1,min(15,round(math.log2(360/span)-1))))
                st.map(mp,latitude='lat',longitude='lon',size=200,color=GREEN,zoom=zoom,height=340,alt='Points approximatifs des quartiers ayant au moins cinq logements visités')
                st.caption('Les repères ci-dessus restent disponibles si les tuiles du fond externe ne se chargent pas.')
            else:st.info('Aucun point ne remplit les conditions de regroupement et de localisation. Le tableau reste disponible.')
    st.download_button('Statistiques agrégées CSV',csv_bytes(stats),'HabitatElles_statistiques.csv','text/csv',on_click='ignore')


def housing_page():
    p=current();rev=st.session_state.rev;title('Logements et campagne','Des repères pseudonymisés, des quartiers et un historique distinct des visites.')
    with st.expander('Cadrer l’enquête'):
        with st.form('campaign_'+str(rev)):
            name=st.text_input('Nom de l’enquête',p['name'],max_chars=120);start=st.date_input('Début de la campagne',date.fromisoformat(p['start']),max_value=date.today())
            sampling=st.selectbox('Méthode de recrutement',SAMPLING,index=SAMPLING.index(p['sampling']));notes=st.text_area('Périmètre et méthode de recrutement',p['sampling_note'],max_chars=900)
            if st.form_submit_button('Enregistrer le cadrage'):
                q=copy.deepcopy(p);q.update(name=name,start=start.isoformat(),sampling=sampling,sampling_note=notes);commit(q)
    with st.expander('Ajouter un quartier',expanded=not p['neighborhoods']):
        with st.form('area_'+str(rev)):
            name=st.text_input('Nom du quartier',max_chars=100);point=st.checkbox('Associer un point approximatif de quartier')
            a,b=st.columns(2);lat=a.number_input('Latitude du point de quartier',-90.,90.,0.,format='%.4f');lon=b.number_input('Longitude du point de quartier',-180.,180.,0.,format='%.4f')
            source=st.text_input('Source et signification du point',max_chars=500)
            st.caption('Le point est facultatif et arrondi à deux décimales. Aucune adresse ou position GPS de logement n’est demandée. Les coordonnées ne sont retenues que si la case est cochée.')
            if st.form_submit_button('Enregistrer le quartier',type='primary'):
                q=copy.deepcopy(p);q['neighborhoods'].append(dict(id=uid(),name=name,lat=round(lat,2) if point else None,lon=round(lon,2) if point else None,source=source if point else ''));commit(q)
    if not p['neighborhoods']:return
    with st.expander('Créer une fiche logement',expanded=not p['houses']):
        with st.form('house_new_'+str(rev)):
            area=st.selectbox('Quartier du nouveau logement',[q['id'] for q in p['neighborhoods']],format_func=lambda i:neighborhood_name(p,i))
            alias=st.text_input('Repère pseudonymisé du logement',max_chars=80,placeholder='Ex. : Logement A. Ne pas saisir de nom de personne ou d’adresse.')
            a,b=st.columns(2);typ=a.selectbox('Type du nouveau logement',HOUSING_TYPES);occupants=b.number_input('Occupants (0 = non renseigné)',0,100,0)
            consent=st.checkbox('Accord pour l’enquête recueilli auprès de l’occupante')
            if st.form_submit_button('Créer le logement',type='primary'):
                q=copy.deepcopy(p);hid=next_house_id(q);q['houses'].append(dict(id=hid,alias=alias,neighborhood_id=area,type=typ,occupants=occupants,consent=consent));st.session_state.active_house=hid;commit(q)
    if not p['houses']:return
    st.dataframe(housing_table(p),hide_index=True,alt='Registre des logements et état de leur dernière visite')
    h=choose_house(p,'housing');pr=priority(p,h['id']);visits=sorted([v for v in p['visits'] if v['house_id']==h['id']],key=lambda v:(v['date'],v['seq']),reverse=True)
    st.subheader(h['id']+' · '+h['alias']);st.write('**'+pr['label']+'**')
    for reason in pr['reasons']:st.caption(reason)
    with st.container(horizontal=True):
        if st.button('Créer une visite pour ce logement',type='primary'):go_house(h['id'],'visites')
        if st.button('Examiner ses besoins'):go_house(h['id'],'besoins')
    if visits:
        st.dataframe(pd.DataFrame([dict(date=v['date'],enqueteur=v['investigator'],ordre=v['seq'],photos=len(v['photos'])) for v in visits]),hide_index=True,alt='Historique des visites de ce logement')
        st.subheader('Photos des deux dernières visites')
        cols=st.columns(2)
        for col,v in zip(cols,visits[:2]):
            with col:
                st.markdown('**'+v['date']+'**')
                if not v['photos']:st.caption('Aucune photographie jointe.')
                for im in v['photos']:st.image(base64.b64decode(im['data']),caption=CATEGORIES[im['aspect']]+' · '+im['caption'],alt='Photo autorisée de défaut, visite du '+v['date'])
        st.caption('Ces images documentent les visites ; aucune résolution de défaut n’est déduite automatiquement de leur comparaison.')
    with st.expander('Corriger ou retirer la fiche'):
        with st.form('house_edit_'+h['id']+'_'+str(rev)):
            alias=st.text_input('Repère corrigé',h['alias'],max_chars=80);typ=st.selectbox('Type corrigé',HOUSING_TYPES,index=HOUSING_TYPES.index(h['type']));occupants=st.number_input('Occupants corrigés (0 = non renseigné)',0,100,h['occupants'])
            if st.form_submit_button('Enregistrer la fiche corrigée'):
                q=copy.deepcopy(p);row=next(x for x in q['houses'] if x['id']==h['id']);row.update(alias=alias,type=typ,occupants=occupants);commit(q)
        sure=st.checkbox('Confirmer le retrait de ce logement, de ses visites, photos et interventions')
        if st.button('Retirer le dossier du logement',disabled=not sure):
            q=copy.deepcopy(p)
            for key in ['houses','visits','tasks']:q[key]=[x for x in q[key] if (x['id'] if key=='houses' else x['house_id'])!=h['id']]
            st.session_state.active_house=None;commit(q)


def visit_page():
    p=current();rev=st.session_state.rev;title('Collecter une visite','Distinguer ce qui est observé, déclaré, mesuré ou encore inconnu.')
    h=choose_house(p,'visit_house')
    if h is None:return
    own=sorted([v for v in p['visits'] if v['house_id']==h['id']],key=lambda v:(v['date'],v['seq']),reverse=True)
    mode=st.segmented_control('Opération de visite',['Nouvelle visite','Corriger une visite'],default='Nouvelle visite')
    if mode=='Corriger une visite':
        if not own:st.info('Aucune visite à corriger pour ce logement.');return
        vid=st.selectbox('Visite à corriger',[v['id'] for v in own],format_func=lambda i:next(v['date']+' · '+v['investigator']+' · ordre '+str(v['seq']) for v in own if v['id']==i));v=next(v for v in own if v['id']==vid)
    else:v=new_visit(p,h['id'])
    fkey=h['id']+'_'+(v['id'] if mode=='Corriger une visite' else 'new')+'_'+str(rev)
    note('L’intensité décrit l’<b>impact d’usage</b> consigné : 1 = limité, 2 = gênant, 3 = important. Ce n’est pas une note de sécurité. Les champs d’intensité ne s’appliquent que si un défaut est présent.')
    with st.form('visit_'+fkey):
        a,b=st.columns(2);when=a.date_input('Date de la visite',date.fromisoformat(v['date']),min_value=date.fromisoformat(p['start']),max_value=date.today());investigator=b.text_input('Code enquêteur',v['investigator'],max_chars=20)
        weather=st.selectbox('Contexte météo de visite',WEATHER,index=WEATHER.index(v['weather']));aspects={}
        for i,(key,label) in enumerate(CATEGORIES.items()):
            old=v['aspects'][key]
            with st.expander(label+' · constat',expanded=i==0):
                state=st.selectbox('État · '+label,STATES,index=STATES.index(old['state']),key=fkey+key+'state')
                origin=st.selectbox('Origine · '+label,ORIGINS,index=ORIGINS.index(old['origin']),key=fkey+key+'origin')
                a,b=st.columns(2);intensity=a.selectbox('Impact d’usage · '+label,[1,2,3],index=max(0,old['intensity']-1),key=fkey+key+'intensity');extent=b.selectbox('Étendue · '+label,EXTENTS,index=EXTENTS.index(old['extent']),key=fkey+key+'extent')
                notes=st.text_area('Pièce et description · '+label,old['notes'],max_chars=600,key=fkey+key+'notes')
                concern=st.checkbox('Signalement préoccupant à faire examiner · '+label,old['concern'],key=fkey+key+'concern')
                reason=st.text_input('Motif du signalement · '+label,old['reason'],max_chars=500,key=fkey+key+'reason')
                aspects[key]=dict(state=state,origin=origin,intensity=intensity if state=='Présent' else 0,extent=extent,notes=notes,concern=concern,reason=reason)
        needs=st.text_area('Besoin prioritaire exprimé par l’occupante',v['needs'],max_chars=800)
        constraints=st.text_area('Contraintes pratiques pour les visites ou travaux',v['constraints'],max_chars=800)
        st.caption('Décrire les difficultés d’usage et d’organisation sans noms, coordonnées personnelles ni informations médicales. Enregistrer une visite incomplète est possible : ses rubriques manquantes resteront visibles.')
        if st.form_submit_button('Enregistrer la visite',type='primary',icon=':material/save:'):
            changed=copy.deepcopy(v);changed.update(date=when.isoformat(),investigator=investigator,weather=weather,aspects=aspects,needs=needs,constraints=constraints);commit(put_visit(p,changed))
    with st.expander('Importer les constats d’une nouvelle visite'):
        st.write('Un fichier décrit les cinq rubriques d’une seule visite. L’import ajoute une visite et conserve l’historique. Le modèle contient des rubriques non observées, sans défaut présumé.')
        st.download_button('Modèle CSV des constats',csv_bytes(template_frame()),'HabitatElles_modele_constats.csv','text/csv',on_click='ignore')
        upload=st.file_uploader('Constats CSV ou Excel',type=['csv','xlsx'],key='aspects_upload_'+str(rev))
        when=st.date_input('Date de la visite importée',date.today(),min_value=date.fromisoformat(p['start']),max_value=date.today())
        investigator=st.text_input('Code enquêteur pour l’import','ENQ-01',max_chars=20)
        needs=st.text_input('Besoin exprimé pour la visite importée',max_chars=800)
        if upload:
            try:
                imported=import_aspects(upload.getvalue(),upload.name);st.dataframe(pd.DataFrame([dict(rubrique=CATEGORIES[k],**a) for k,a in imported.items()]),hide_index=True,alt='Aperçu des constats importés')
                sure=st.checkbox('Créer une nouvelle visite avec ces constats vérifiés')
                if st.button('Appliquer la visite importée',disabled=not sure):
                    nv=new_visit(p,h['id']);nv.update(date=when.isoformat(),investigator=investigator,needs=needs,aspects=imported);commit(put_visit(p,nv))
            except Exception as ex:st.error('Import non appliqué : '+str(ex))
    if not own:return
    with st.expander('Photos autorisées : ajouter ou retirer'):
        vid=st.selectbox('Visite associée à la photo',[v['id'] for v in own],format_func=lambda i:next(v['date']+' · ordre '+str(v['seq']) for v in own if v['id']==i));pv=next(v for v in own if v['id']==vid)
        st.caption('Photographier les défauts seuls, sans personnes ni documents identifiants. Les images sont réduites et leurs métadonnées EXIF supprimées ; aucun visage n’est détecté ou masqué automatiquement. Deux photos par visite, quarante par enquête.')
        method=st.segmented_control('Source de la photographie',['Importer une photo','Prendre une photo'],default='Importer une photo')
        file=st.camera_input('Prendre la photographie du défaut',key='camera_'+vid+'_'+str(rev)) if method=='Prendre une photo' else st.file_uploader('Photo du défaut (JPEG ou PNG)',type=['jpg','jpeg','png'],key='photo_'+vid+'_'+str(rev))
        aspect=st.selectbox('Rubrique de la photographie',list(CATEGORIES),format_func=lambda k:CATEGORIES[k]);caption=st.text_input('Légende de la photographie',max_chars=200)
        consent=st.checkbox('Accord explicite recueilli pour cette photographie',value=pv['photo_consent'])
        if st.button('Ajouter la photographie',disabled=file is None):
            if not consent:raise ValueError('Recueillir l’accord photographique avant d’ajouter une image.')
            q=copy.deepcopy(p);qv=next(v for v in q['visits'] if v['id']==vid);qv['photo_consent']=True;qv['photos'].append(dict(id=uid(),caption=caption,aspect=aspect,data=normalize_photo(file.getvalue())));commit(q)
        for im in pv['photos']:st.image(base64.b64decode(im['data']),caption=im['caption'],alt='Photo autorisée : '+CATEGORIES[im['aspect']])
        if pv['photos']:
            photo_id=st.selectbox('Photographie à retirer',[im['id'] for im in pv['photos']],format_func=lambda i:next(im['caption'] for im in pv['photos'] if im['id']==i));sure=st.checkbox('Confirmer le retrait de cette photographie')
            if st.button('Retirer la photographie',disabled=not sure):
                q=copy.deepcopy(p);qv=next(v for v in q['visits'] if v['id']==vid);qv['photos']=[im for im in qv['photos'] if im['id']!=photo_id];qv['photo_consent']=bool(qv['photos']);commit(q)


def needs_page():
    p=current();title('Besoins de réhabilitation','Une orientation expliquée, à examiner avec les constats et les besoins exprimés.')
    if not p['houses']:st.info('Aucun logement enregistré.');return
    table=housing_table(p);level=st.selectbox('Filtrer les orientations',['Toutes']+PRIORITIES)
    shown=table if level=='Toutes' else table[table.priorite==level]
    st.dataframe(shown,hide_index=True,alt='Logements et orientations proposées, données manquantes visibles')
    h=choose_house(p,'needs_house');v=latest_visits(p).get(h['id']);pr=priority(p,h['id'],v)
    st.subheader(h['id']+' · '+pr['label'])
    for reason in pr['reasons']:st.write('• '+reason)
    if v:
        st.caption('Dernière visite : '+v['date']+' · '+v['investigator'])
        st.dataframe(pd.DataFrame([dict(rubrique=CATEGORIES[k],etat=a['state'],origine=a['origin'],impact_usage=a['intensity'] if a['state']=='Présent' else None,description=a['notes']) for k,a in v['aspects'].items()]),hide_index=True,alt='Constats de la dernière visite et origine des informations')
        note('<b>Priorité exprimée :</b> '+html.escape(v['needs'] or 'Non renseignée')+'<br><b>Contraintes pratiques :</b> '+html.escape(v['constraints'] or 'Non renseignées'))
        with st.container(horizontal=True):
            if st.button('Préparer une intervention pour ce logement',type='primary'):go_house(h['id'],'suivi')
            if st.button('Compléter ou corriger une visite'):go_house(h['id'],'visites')
    with st.expander('Comment l’orientation est-elle proposée ?'):
        st.write('1. Un signalement préoccupant non clôturé entraîne une visite technique à prioriser. 2. Sinon, un défaut avec impact d’usage de 2 ou 3 conduit à programmer une réhabilitation. 3. Un défaut d’impact 1 conduit à organiser l’entretien. 4. En l’absence de défaut consigné, les données manquantes conduisent à compléter les informations ; cinq rubriques évaluées sans défaut conduisent à une surveillance courante.')
        st.caption('Une information nouvelle manquante n’efface pas un ancien signalement préoccupant. Seule une intervention reliée à la visite et à la rubrique, vérifiée avec une source et un résultat de résolution, clôture ce signalement. Ce classement ne diagnostique pas la structure et ne dépend pas du nombre d’occupants ou de la situation familiale.')
    with st.expander('Sensibilité à la règle d’impact d’usage'):
        st.write('Comparer le seuil 2 au seuil 3 montre les dossiers dont le classement réhabilitation/entretien dépend de ce choix pédagogique. Les signalements préoccupants restent prioritaires dans les deux cas.')
        other=housing_table(p,cutoff=3);sensitivity=table[['logement','priorite']].rename(columns={'priorite':'seuil_2'}).merge(other[['logement','priorite']].rename(columns={'priorite':'seuil_3'}),on='logement');sensitivity['change']=sensitivity.seuil_2!=sensitivity.seuil_3
        st.dataframe(sensitivity,hide_index=True,alt='Comparaison de la règle d’impact au seuil 2 et au seuil 3')


def new_task(hid,vid):
    return dict(id=uid(),house_id=hid,visit_id=vid,aspect='toiture',stage=STAGES[0],action='',hypothesis='',checks='',planned_date=None,work_date=None,verified_date=None,result=RESULTS[0],verification_note='',verification_source='',budget_include=False,cost=None,cost_source='')


def tracking_page():
    p=current();rev=st.session_state.rev;title('Interventions et suivi','Un travail réalisé et un problème résolu sont deux informations distinctes.')
    h=choose_house(p,'task_house')
    if h is None:return
    visits=sorted([v for v in p['visits'] if v['house_id']==h['id']],key=lambda v:(v['date'],v['seq']),reverse=True)
    if not visits:st.info('Enregistrez une visite avant de préparer une intervention.');return
    own=[t for t in p['tasks'] if t['house_id']==h['id']]
    if own:st.dataframe(pd.DataFrame([dict(action=t['action'],rubrique=CATEGORIES[t['aspect']],stade=t['stage'],date_cible=t['planned_date'],resultat=t['result']) for t in own]),hide_index=True,alt='Interventions de ce logement et étapes de suivi')
    mode=st.segmented_control('Opération d’intervention',['Nouvelle intervention','Corriger une intervention'],default='Nouvelle intervention')
    if mode=='Corriger une intervention':
        if not own:st.info('Aucune intervention à corriger.');return
        tid=st.selectbox('Intervention à corriger',[t['id'] for t in own],format_func=lambda i:next(CATEGORIES[t['aspect']]+' · '+t['action'][:65] for t in own if t['id']==i));t=next(t for t in own if t['id']==tid)
    else:t=new_task(h['id'],visits[0]['id'])
    with st.form('task_'+h['id']+'_'+(t['id'] if mode=='Corriger une intervention' else 'new')+'_'+str(rev)):
        ids=[v['id'] for v in visits];vid=st.selectbox('Visite de référence de l’intervention',ids,index=ids.index(t['visit_id']),format_func=lambda i:next(v['date']+' · ordre '+str(v['seq']) for v in visits if v['id']==i))
        aspect=st.selectbox('Rubrique à traiter',list(CATEGORIES),index=list(CATEGORIES).index(t['aspect']),format_func=lambda k:CATEGORIES[k])
        action=st.text_area('Action proposée ou réalisée',t['action'],max_chars=600)
        hypothesis=st.text_area('Cause supposée, à vérifier',t['hypothesis'],max_chars=600);checks=st.text_area('Vérifications nécessaires avant décision',t['checks'],max_chars=600)
        stage=st.selectbox('Étape de suivi',STAGES,index=STAGES.index(t['stage']))
        planned=st.checkbox('Date cible connue',t['planned_date'] is not None);pdue=st.date_input('Date cible',date.fromisoformat(t['planned_date']) if t['planned_date'] else date.today()+timedelta(days=14))
        worked=st.checkbox('Une date de travaux réalisés est disponible',t['work_date'] is not None);work=st.date_input('Date des travaux réalisés',date.fromisoformat(t['work_date']) if t['work_date'] else date.today(),max_value=date.today())
        result=st.selectbox('Résultat de la vérification',RESULTS,index=RESULTS.index(t['result']))
        verified=st.date_input('Date de vérification du résultat',date.fromisoformat(t['verified_date']) if t['verified_date'] else date.today(),max_value=date.today())
        verification_source=st.text_input('Source de la vérification',t['verification_source'],max_chars=400)
        verification_note=st.text_area('Note de vérification du résultat',t['verification_note'],max_chars=600)
        st.caption('Au stade Résultat vérifié : date, résultat, source et note sont requis. Une résolution ne clôture que le signalement associé à cette visite et cette rubrique. Une visite technique peut conclure sans travaux ; sa conclusion doit être documentée.')
        with st.expander('Estimation et scénario de budget'):
            known=st.checkbox('Une estimation de coût est disponible',t['cost'] is not None);old=t['cost'] or dict(law='Triangulaire',min=0.,mode=0.,max=0.)
            law=st.selectbox('Loi du coût estimé',LAWS,index=LAWS.index(old['law']));a,b,c=st.columns(3)
            low=a.number_input('Coût minimum (FCFA)',0.,1e9,float(old['min']),step=1000.);central=b.number_input('Coût central / mode (FCFA)',0.,1e9,float(old['mode']),step=1000.);high=c.number_input('Coût maximum (FCFA)',0.,1e9,float(old['max']),step=1000.)
            cost_source=st.text_input('Source du coût ou justification de la fourchette',t['cost_source'],max_chars=500)
            include=st.checkbox('Inclure dans le scénario de dépenses futures',t['budget_include'])
            st.caption('Fixe : seule la centrale est retenue. Uniforme : centrale recalculée au milieu. Triangulaire : centrale = mode. Sans estimation, le coût reste non renseigné ; il n’est pas assimilé à zéro. Les travaux réalisés sont exclus du scénario futur.')
        if st.form_submit_button('Enregistrer l’intervention',type='primary'):
            if law=='Fixe':low=high=central
            elif law=='Uniforme':central=(low+high)/2
            cost=dict(law=law,min=low,mode=central,max=high) if known else None
            changed=dict(t,visit_id=vid,aspect=aspect,action=action,hypothesis=hypothesis,checks=checks,stage=stage,planned_date=pdue.isoformat() if planned else None,
                         work_date=work.isoformat() if worked and stage in STAGES[-2:] else None,verified_date=verified.isoformat() if stage==STAGES[-1] else None,
                         result=result if stage==STAGES[-1] else RESULTS[0],verification_source=verification_source,verification_note=verification_note,
                         cost=cost,cost_source=cost_source if known else '',budget_include=include)
            q=copy.deepcopy(p);q['tasks']=[x for x in q['tasks'] if x['id']!=t['id']]+[changed];commit(q)
    if mode=='Corriger une intervention':
        with st.expander('Retirer cette intervention'):
            sure=st.checkbox('Confirmer le retrait de cette intervention')
            if st.button('Retirer l’intervention',disabled=not sure):
                q=copy.deepcopy(p);q['tasks']=[x for x in q['tasks'] if x['id']!=t['id']];commit(q)
    st.subheader('Scénario budgétaire de l’enquête entière')
    selected,missing=budget_scope(p)
    st.caption(f'{len(selected)} intervention(s) future(s) chiffrée(s) sélectionnée(s) ; {missing} sélectionnée(s) sans coût. Les autres interventions et les travaux déjà réalisés sont exclus. Les coûts proviennent uniquement de vos estimations.')
    with st.form('budget_sim'):
        a,b=st.columns(2);count=a.selectbox('Tirages du scénario budgétaire',[1000,10000,50000,100000],index=1);seed=b.number_input('Graine du budget',0,2**32-1,42)
        if st.form_submit_button('Simuler le sous-total chiffré',type='primary'):
            costs=tuple((t['cost']['law'],t['cost']['min'],t['cost']['mode'],t['cost']['max']) for t in selected)
            values=budget_simulation(costs,count,seed);st.session_state.budget_result=dict(values=values,count=count,seed=seed,fingerprint=hashlib.sha256(backup_bytes(p)).hexdigest(),missing=missing)
    s=st.session_state.get('budget_result')
    if s and s['fingerprint']==hashlib.sha256(backup_bytes(p)).hexdigest():
        x=s['values'];budget=st.number_input('Budget disponible pour ce sous-total (FCFA)',0.,1e12,250000.,step=10000.)
        prob=float((x<=budget).mean())
        kpis([('Sous-total moyen',fmt(x.mean(),0)+' FCFA','Selon les coûts saisis'),('P95 du sous-total',fmt(np.quantile(x,.95),0)+' FCFA','95 % des valeurs simulées en dessous'),('Budget suffisant',fmt(100*prob)+' %','P(sous-total ≤ budget)'),('Non chiffrées',missing,'Interventions sélectionnées exclues du calcul')])
        note('Cette probabilité concerne uniquement le <b>sous-total chiffré sélectionné</b>, pas la couverture de tous les besoins de réhabilitation. Les coûts sont supposés indépendants ; leurs fourchettes et lois sont des hypothèses.')
        fig=go.Figure(go.Histogram(x=x,nbinsx=40,histnorm='probability',marker_color=GREEN));fig.add_vline(x=budget,line_color=CLAY,line_dash='dash');fig.update_xaxes(title='Sous-total simulé · FCFA');fig.update_yaxes(title='Fréquence');plot(fig,'Distribution du sous-total d’interventions futures et budget disponible')
        st.caption(f"{s['count']} tirages ; graine {s['seed']}. Toute modification enregistrée invalide ce scénario. P05–P95 est un intervalle de résultats sous hypothèses, pas un intervalle de confiance de la moyenne.")
        st.download_button('Tirages du budget CSV',csv_bytes(pd.DataFrame(dict(sous_total_FCFA=x))),'HabitatElles_budget.csv','text/csv',on_click='ignore')
        meta=dict(app="Habitat'Elles · budget",count=s['count'],seed=s['seed'],budget=budget,probability=prob,non_chiffrees=missing,costs=[dict(intervention=t['id'],**t['cost']) for t in selected])
        st.download_button('Hypothèses budgétaires JSON',json.dumps(meta,ensure_ascii=False,indent=2).encode(),'HabitatElles_budget_hypotheses.json','application/json',on_click='ignore')


def reports_page():
    p=current();title('Rapports et sauvegardes','Partager une synthèse agrégée et conserver séparément les données détaillées de l’enquête.')
    prefix='FICTIF_' if p['demo'] else ''
    note('Le projet reste dans <b>votre session</b>. Téléchargez la sauvegarde JSON avant de fermer ou recharger la connexion. Le prototype collecte via un navigateur connecté ; il n’est pas une application de terrain entièrement hors ligne.')
    with st.container(horizontal=True):
        st.download_button('Rapport agrégé PDF',pdf_bytes(p),prefix+'HabitatElles_synthese.pdf','application/pdf',type='primary',on_click='ignore')
        st.download_button('Statistiques CSV',csv_bytes(rates(p)),prefix+'HabitatElles_statistiques.csv','text/csv',on_click='ignore')
    st.caption('Ces deux exports portent sur le registre entier. Le PDF ne contient ni fiches individuelles ni photos ; les quartiers de moins de cinq logements sont regroupés. Les filtres du tableau de bord ne s’appliquent pas à ces exports.')
    with st.expander('Données détaillées et fiche de suivi'):
        st.caption('Le classeur et la sauvegarde sont des fichiers détaillés à conserver selon les accords d’enquête. Les photographies restent uniquement dans le JSON complet et dans une fiche PDF sur demande.')
        with st.container(horizontal=True):
            st.download_button('Classeur enquête détaillé',excel_bytes(p),prefix+'HabitatElles_enquete.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',on_click='ignore')
            st.download_button('Sauvegarde complète JSON (photos incluses)',backup_bytes(p),prefix+'HabitatElles_sauvegarde.json','application/json',on_click='ignore')
            st.download_button('Constats détaillés CSV',csv_bytes(aspect_rows(p)),prefix+'HabitatElles_constats.csv','text/csv',on_click='ignore')
        if p['houses']:
            h=choose_house(p,'report_house');include=st.checkbox('Inclure les photographies autorisées dans la fiche PDF')
            st.download_button('Fiche de suivi PDF (détaillée)',pdf_bytes(p,h['id'],include),prefix+h['id']+'_suivi.pdf','application/pdf',on_click='ignore')
    with st.expander('Restaurer une sauvegarde'):
        file=st.file_uploader("Sauvegarde Habitat’Elles JSON",type=['json'])
        if file:
            try:
                q=restore(file.getvalue());st.write(q['name']+f" · {len(q['houses'])} logements · {len(q['visits'])} visites · "+('FICTIF' if q['demo'] else 'PERSONNEL'))
                sure=st.checkbox('Remplacer l’espace correspondant par cette sauvegarde')
                if st.button('Restaurer l’enquête',disabled=not sure):st.session_state.demo=q['demo'];st.session_state.active_house=None;commit(q)
            except Exception as ex:st.error('Restauration non appliquée : '+str(ex))
    with st.expander('Démarrer ou réinitialiser un espace'):
        if st.button('Réinitialiser la démonstration fictive'):st.session_state.demo=True;st.session_state.active_house=None;commit(example())
        sure=st.checkbox('J’ai sauvegardé l’étude personnelle à conserver')
        if st.button('Créer une enquête personnelle vide',disabled=not sure):st.session_state.demo=False;st.session_state.active_house=None;commit(new_project())
    st.subheader('Tests guidés et lecture des résultats')
    tests=[
        ('1 · Logements et visites','L’exemple contient 30 logements et 32 visites. Deux logements ont une visite ancienne supplémentaire. Les statistiques prennent la plus récente de chaque logement : 30 unités, sans double comptage.'),
        ('2 · Toitures et dénominateur','Dans l’exemple : 12 toitures avec défaut, 8 sans défaut, 5 non observées et 5 inconnues. Fréquence parmi les 20 toitures évaluées : 12/20 = 60 %. Les dix autres restent visibles. On ne présente pas 12/30 comme la fréquence parmi les toitures évaluées.'),
        ('3 · Priorité expliquée','HE-0008 contient un signalement préoccupant fictif sur les murs. Il reste orienté vers une visite technique tant que le signalement n’a pas de vérification documentée de résolution reliée à la même visite et rubrique. Une nouvelle visite inconnue ne l’efface pas.'),
        ('4 · Comparaison de règles','Dans Besoins, comparer les seuils d’impact 2 et 3. Les dossiers dépendant de ce choix sont signalés ; le signalement préoccupant de HE-0008 reste prioritaire. Aucun revenu ou statut familial n’entre dans cette règle.'),
        ('5 · Budget vérifiable','Le scénario fictif sélectionne deux interventions : la première uniforme de 100000 à 200000 FCFA, la deuxième fixe à 100000. Sous-total de 200000 à 300000 ; moyenne 250000 ; P05=205000 ; P95=295000. P(sous-total≤250000)=50 % théoriquement. À 10000 tirages, variation approximative de ±0,98 point dans 95 % des répétitions, sans garantie pour chaque graine.'),
        ('6 · Collecte et photos','Enregistrer une visite avec une rubrique non observée : elle reste manquante. Ajouter une photo exige l’accord correspondant et une légende. L’image est redimensionnée et les EXIF supprimés. Les photos ne figurent pas dans la synthèse agrégée.'),
        ('7 · Historique et résolution','Marquer Travaux réalisés ne signifie pas Résultat vérifié. Pour ce dernier stade, une date, une conclusion, une note et une source sont requises. Un problème persistant maintient un signalement préoccupant non résolu.'),
        ('8 · Sauvegardes et limites','Le JSON retrouve quartiers, logements, visites, interventions et photos. La restauration remplace l’espace fictif ou personnel correspondant après confirmation, en conservant l’autre. Les photos, libellés et accords restent à examiner avant tout partage de fichiers détaillés.')]
    for label,body in tests:
        with st.expander(label):st.write(body)
    st.caption('Les constats et orientations de ce prototype sont pédagogiques. Une inspection compétente reste nécessaire pour caractériser un problème technique. Les intervalles statistiques ne corrigent pas les biais de sélection. La prise de photo dépend des permissions du navigateur et ne remplace pas une collecte hors ligne.')
    for label,url in REFS:st.markdown(f'[{label}]({url})')


def main():
    st.set_page_config(page_title="Habitat’Elles · Constats et réhabilitation",page_icon=':material/cottage:',layout='wide');st.html(CSS)
    ss=st.session_state
    if 'real_project' not in ss:ss.update(real_project=new_project(),demo_project=example(),demo=True,rev=0,flash='',active_house=None,filters={},budget_result=None)
    p=current();st.html('<div class="mast"><div class="brand">Habitat<span>’Elles</span></div><span class="tag">'+('EXEMPLE FICTIF' if ss.demo else 'ENQUÊTE PERSONNELLE')+' / COLLECTE CONNECTÉE</span></div>')
    routes=[('accueil',dashboard,'Tableau de bord','dashboard'),('logements',housing_page,'Logements','home'),('visites',visit_page,'Visites','assignment'),('besoins',needs_page,'Besoins','format_list_bulleted'),('suivi',tracking_page,'Interventions','construction'),('rapports',reports_page,'Rapports','description')]
    pages={path:st.Page(func,title=label,icon=':material/'+icon+':',url_path=path,default=path=='accueil') for path,func,label,icon in routes};ss.pages=pages
    with st.sidebar:st.html('<div class="sidebrand">Habitat’Elles</div><div class="sidenote">Collecter avec attention.<br>Suivre les besoins dans le temps.</div>')
    nav=st.navigation(list(pages.values()))
    with st.sidebar:
        st.caption('ESPACE DE TRAVAIL');st.write('**'+p['name']+'**')
        if ss.demo:
            if st.button('Revenir à mon enquête',icon=':material/folder:'):ss.demo=False;ss.active_house=None;ss.budget_result=None;ss.rev+=1;st.rerun()
        else:
            if st.button('Ouvrir l’exemple fictif',icon=':material/science:'):ss.demo=True;ss.active_house=None;ss.budget_result=None;ss.rev+=1;st.rerun()
        st.caption('Version 1.0 · Données propres à la session')
    if ss.flash:st.toast(ss.flash,icon=':material/check_circle:');ss.flash=''
    try:nav.run()
    except (ValueError,TypeError,KeyError,OverflowError) as ex:st.error('Action non appliquée : '+str(ex))
    st.html('<div class="foot">Habitat’Elles · Prototype pédagogique de collecte et de suivi. Les orientations proposées restent à examiner avec les occupantes et les intervenants compétents.</div>')


if __name__=='__main__':main()
